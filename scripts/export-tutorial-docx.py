import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TABLE_WIDTH_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(index, len(widths) - 1)])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in [
        ("Title", 22, RGBColor(31, 78, 121)),
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(54, 96, 146)),
        ("Heading 3", 11.5, RGBColor(89, 89, 89)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header
    p = header.paragraphs[0]
    p.text = "FrameNotes Tutorial"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(128, 128, 128)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Page ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def add_table(doc, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    widths = [int(TABLE_WIDTH_DXA / columns)] * columns
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    set_table_width(table, widths)

    for r, row in enumerate(rows):
        for c in range(columns):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if r == 0:
                run.bold = True
                set_cell_shading(cell, "EAF2F8")
    doc.add_paragraph()


def add_image(doc, md_path, image_ref):
    image_path = (md_path.parent / image_ref).resolve()
    if not image_path.exists():
        doc.add_paragraph(f"[Missing image: {image_ref}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.1))


def parse_markdown(md_path):
    blocks = []
    table_rows = []
    in_code = False
    code_lines = []

    for raw_line in md_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code:
                blocks.append(("code", "\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            table_rows.append(cells)
            continue
        if table_rows:
            blocks.append(("table", table_rows))
            table_rows = []

        image_match = re.match(r"!\[[^\]]*\]\((.+)\)", line)
        if image_match:
            blocks.append(("image", image_match.group(1)))
        elif line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h1", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h2", line[4:].strip()))
        elif line.startswith("- [ ] "):
            blocks.append(("checkbox", line[6:].strip()))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        elif re.match(r"^\d+\. ", line):
            blocks.append(("number", re.sub(r"^\d+\. ", "", line).strip()))
        elif line:
            blocks.append(("para", line))
        else:
            blocks.append(("blank", ""))

    if table_rows:
        blocks.append(("table", table_rows))
    return blocks


def add_blocks(doc, md_path, blocks):
    for kind, content in blocks:
        if kind == "title":
            doc.add_paragraph(content, style="Title")
        elif kind == "h1":
            doc.add_paragraph(content, style="Heading 1")
        elif kind == "h2":
            doc.add_paragraph(content, style="Heading 2")
        elif kind == "para":
            doc.add_paragraph(content)
        elif kind == "bullet":
            doc.add_paragraph(content, style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(content, style="List Number")
        elif kind == "checkbox":
            doc.add_paragraph(f"☐ {content}", style="List Bullet")
        elif kind == "table":
            add_table(doc, content)
        elif kind == "image":
            add_image(doc, md_path, content)
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9)
        elif kind == "blank":
            continue


def main():
    parser = argparse.ArgumentParser(description="Export final tutorial Markdown to DOCX.")
    parser.add_argument("markdown", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--name-prefix")
    args = parser.parse_args()

    md_path = args.markdown.resolve()
    if args.output:
        output = args.output.resolve()
    elif args.name_prefix:
        output = md_path.with_name(f"{args.name_prefix}.docx")
    else:
        output = md_path.with_suffix(".docx")

    doc = Document()
    style_document(doc)
    add_blocks(doc, md_path, parse_markdown(md_path))
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
